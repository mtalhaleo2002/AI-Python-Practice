import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def create_document():
    doc = docx.Document()

    # Page setup - Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Title Banner / Header Table
    title_table = doc.add_table(rows=1, cols=1)
    title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = title_table.cell(0, 0)
    set_cell_background(cell, "003366") # Deep Navy

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("DigiSkills 3.0 Batch-04\nArtificial Intelligence using Python")
    run.font.name = "Arial"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(12)
    run2 = p2.add_run("Hands-on Exercise No. 01 Solution")
    run2.font.name = "Arial"
    run2.font.size = Pt(13)
    run2.font.color.rgb = RGBColor(200, 220, 240)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Helper function for adding task sections
    def add_task_section(task_num, task_title, description_text, code_file, png_file):
        # Heading 1
        h1 = doc.add_paragraph()
        h1.paragraph_format.space_before = Pt(18)
        h1.paragraph_format.space_after = Pt(6)
        r = h1.add_run(f"Task {task_num}: {task_title}")
        r.font.name = "Arial"
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 51, 102)

        # Description
        p_desc = doc.add_paragraph(description_text)
        p_desc.paragraph_format.space_after = Pt(8)
        for run in p_desc.runs:
            run.font.name = "Arial"
            run.font.size = Pt(11)

        # Subheading for Code
        h_code = doc.add_paragraph()
        h_code.paragraph_format.space_before = Pt(6)
        h_code.paragraph_format.space_after = Pt(4)
        r_code = h_code.add_run("Python Code:")
        r_code.font.name = "Arial"
        r_code.font.size = Pt(11)
        r_code.font.bold = True
        r_code.font.color.rgb = RGBColor(80, 80, 80)

        # Code box (Table with 1 cell)
        with open(code_file, 'r') as f:
            code_text = f.read()

        code_table = doc.add_table(rows=1, cols=1)
        code_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        c_cell = code_table.cell(0, 0)
        set_cell_background(c_cell, "F4F4F4")
        cp = c_cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(6)
        cp.paragraph_format.space_after = Pt(6)
        cp.paragraph_format.left_indent = Inches(0.1)
        cp.paragraph_format.right_indent = Inches(0.1)
        c_run = cp.add_run(code_text)
        c_run.font.name = "Courier New"
        c_run.font.size = Pt(9.5)
        c_run.font.color.rgb = RGBColor(50, 50, 50)

        # Subheading for Output & Screenshot
        h_out = doc.add_paragraph()
        h_out.paragraph_format.space_before = Pt(12)
        h_out.paragraph_format.space_after = Pt(4)
        r_out = h_out.add_run("Code Output & Execution Screenshot:")
        r_out.font.name = "Arial"
        r_out.font.size = Pt(11)
        r_out.font.bold = True
        r_out.font.color.rgb = RGBColor(80, 80, 80)

        # Insert Screenshot Image
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(18)
        run_img = p_img.add_run()
        run_img.add_picture(png_file, width=Inches(6.2))

    # Task 1
    add_task_section(
        1,
        "Mobile Phone Shop Sales & Tax Calculation",
        "This program takes daily sales amounts for 5 sales employees using a for loop, calculates tax according to the specified tiered rules (< 50,000: 5%, 50,000 - 100,000: 10%, > 100,000: 15%), displays individual sales and tax amounts, and outputs the total and average sales.",
        "task1.py",
        "output_solution/task1.png"
    )

    # Task 2
    add_task_section(
        2,
        "NumPy Array Operations & Type Conversion",
        "This program utilizes the NumPy library to create an array of student marks [78, 85, 92, 67, 88], displays the original array along with its integer data type, converts the array into float data type using .astype(float), and displays the converted array.",
        "task2.py",
        "output_solution/task2.png"
    )

    # Task 3
    add_task_section(
        3,
        "Bank Account System using OOP",
        "This program implements a Bank Account System using Object-Oriented Programming (OOP) with Customer and Account classes. It demonstrates customer creation, account initialization with opening balance, money deposits, valid and invalid money withdrawals, and balance inquiries.",
        "task3.py",
        "output_solution/task3.png"
    )

    # Save document
    doc.save("output_solution/DSTP3.0_Batch-04_AIP301_Exercise1_Solution.docx")
    print("Word document generated successfully at output_solution/DSTP3.0_Batch-04_AIP301_Exercise1_Solution.docx")

if __name__ == "__main__":
    create_document()
